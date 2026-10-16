"""导入管理业务服务

提供：
- ImportTaskService: 任务查询
- ImportLogService: 日志查询
- ExcelImportService: Excel 导入主流程（上传、解析、导入、错误导出）

依据 DATABASE.md 7.9 Excel 导入流程：
1. 创建 import_task (academic_year_id, role, status=PROCESSING, file_name, ...)
2. 解析 Excel 文件（Parser 层）
3. 逐行处理（Importer 层）：校验 → 创建 user+profile → 记录 import_log
4. 更新 import_task SET status=SUCCESS/PARTIAL/FAILED, success_count=?, failed_count=?, finished_at=NOW()

业务规则（PROJECT.md 7.13）：
- 单文件大小限制 5MB
- 单文件行数限制 5000
- 单行失败不阻塞整批
- 重复 username 跳过该行，标记为失败
- 导入异步处理，立即返回任务 ID（当前阶段同步处理，Phase D 改为异步）
"""
from __future__ import annotations

import time
from io import BytesIO

from django.utils import timezone

from apps.common.enums import ImportTaskStatus, UserRole
from apps.common.error_codes import ErrorCode
from apps.common.exceptions import BusinessException
from apps.common.logger import get_logger
from apps.common.services import BaseService
from apps.imports.excel.importer import BatchImportResult, ExcelImporter
from apps.imports.excel.parser import ExcelParser
from apps.imports.excel.template import ExcelTemplateGenerator
from apps.imports.models import ImportLog, ImportTask
from apps.imports.repositories import (
    ImportLogRepository,
    ImportTaskRepository,
)
from apps.monitoring.services import MetricsService

logger = get_logger("imports.services")


class ImportTaskService(BaseService):
    """导入任务业务服务"""

    repository = ImportTaskRepository()

    def list(self, filters: dict | None = None, order_by: list | None = None):
        """查询导入任务列表，默认按 ID 倒序"""
        return self.repository.list(filters=filters, order_by=order_by or ["-id"])

    def get_by_id(self, pk: int) -> ImportTask:
        return self.repository.get_by_id(pk)

    def list_by_academic_year(self, academic_year_id: int):
        """按学年查询导入任务"""
        return self.repository.list_by_academic_year(academic_year_id)

    def list_by_status(self, status: int):
        """按状态查询导入任务"""
        return self.repository.list_by_status(status)

    def list_by_role(self, role: int):
        """按角色查询导入任务"""
        return self.repository.list_by_role(role)


class ImportLogService(BaseService):
    """导入日志业务服务（只读）"""

    repository = ImportLogRepository()

    def list(self, filters: dict | None = None, order_by: list | None = None):
        """查询导入日志列表，默认按 ID 倒序"""
        return self.repository.list(filters=filters, order_by=order_by or ["-id"])

    def get_by_id(self, pk: int) -> ImportLog:
        return self.repository.get_by_id(pk)

    def list_by_task(self, task_id: int):
        """按任务查询导入日志"""
        return self.repository.list_by_task(task_id)

    def list_by_task_and_status(self, task_id: int, status: int):
        """按任务+状态查询导入日志"""
        return self.repository.list_by_task_and_status(task_id, status)

    def list_failures(self, task_id: int):
        """查询任务的失败日志"""
        return self.repository.list_failures(task_id)


class ExcelImportService:
    """Excel 导入业务服务

    主流程：
    1. 创建 import_task (status=PROCESSING)
    2. 解析 Excel（Parser）
    3. 逐行导入（Importer）
    4. 更新 import_task 统计和状态
    5. 返回任务实例

    当前阶段同步处理；Phase D 接入 Celery 后改为异步。
    """

    task_repository = ImportTaskRepository()
    log_repository = ImportLogRepository()
    parser = ExcelParser()
    importer = ExcelImporter()

    @classmethod
    def upload_and_import(
        cls,
        file_content: bytes,
        file_name: str,
        file_size: int,
        role: int,
        academic_year_id: int,
        triggered_by,
    ) -> ImportTask:
        """上传 Excel 并执行导入

        Args:
            file_content: Excel 文件二进制内容
            file_name: 原始文件名
            file_size: 文件大小（字节）
            role: UserRole 枚举值
            academic_year_id: 学年 ID
            triggered_by: 触发管理员 User 实例

        Returns:
            ImportTask 实例（含最终状态和统计）
        """
        total_start = time.perf_counter()
        try:
            # 1. 校验角色
            if role not in (
                UserRole.STUDENT.value,
                UserRole.TEACHER.value,
                UserRole.ADMIN.value,
            ):
                raise BusinessException(
                    f"不支持的角色: {role}",
                    code=ErrorCode.PARAM_INVALID,
                )

            # 2. 创建 import_task (status=PROCESSING) — upload 阶段
            upload_start = time.perf_counter()
            task = cls.task_repository.create(
                academic_year_id=academic_year_id,
                role=role,
                status=ImportTaskStatus.PROCESSING.value,
                file_name=file_name,
                file_size=file_size,
                total_rows=0,
                success_count=0,
                failed_count=0,
                started_at=timezone.now(),
                triggered_by=triggered_by,
            )
            MetricsService.record_excel_import(
                "upload", time.perf_counter() - upload_start
            )
            logger.info(
                "Import task created: id=%s, role=%s, file=%s",
                task.id, role, file_name,
            )

            try:
                # 3. 解析 Excel — parse 阶段
                parse_start = time.perf_counter()
                rows = cls.parser.parse(file_content, role)
                MetricsService.record_excel_import(
                    "parse", time.perf_counter() - parse_start
                )
                cls.task_repository.update(task, {"total_rows": len(rows)})

                # 4. 逐行导入 — write 阶段（含逐行校验，importer 内部记录 validate 耗时）
                write_start = time.perf_counter()
                batch_result: BatchImportResult = cls.importer.import_rows(
                    task, rows, role
                )
                MetricsService.record_excel_import(
                    "write", time.perf_counter() - write_start
                )
                # importer 内部累积的 validate 阶段耗时
                if batch_result.validate_duration > 0:
                    MetricsService.record_excel_import(
                        "validate", batch_result.validate_duration
                    )

                # 5. 更新任务状态和统计
                final_status = cls._compute_final_status(batch_result)
                cls.task_repository.update(
                    task,
                    {
                        "status": final_status,
                        "success_count": batch_result.success_count,
                        "failed_count": batch_result.failed_count,
                        "finished_at": timezone.now(),
                    },
                )

                logger.info(
                    "Import task finished: id=%s, status=%s, success=%s, failed=%s",
                    task.id, final_status,
                    batch_result.success_count, batch_result.failed_count,
                )

                # 重新获取任务（含最新状态）
                return cls.task_repository.get_by_id(task.id)

            except BusinessException as exc:
                # 解析失败等业务异常
                cls.task_repository.update(
                    task,
                    {
                        "status": ImportTaskStatus.FAILED.value,
                        "error_message": str(exc.detail)[:512],
                        "finished_at": timezone.now(),
                    },
                )
                logger.error(
                    "Import task failed: id=%s, error=%s", task.id, exc.detail
                )
                raise

            except Exception as exc:
                # 未预期异常
                cls.task_repository.update(
                    task,
                    {
                        "status": ImportTaskStatus.FAILED.value,
                        "error_message": str(exc)[:512],
                        "finished_at": timezone.now(),
                    },
                )
                logger.exception(
                    "Import task unexpected error: id=%s", task.id
                )
                raise BusinessException(
                    f"导入处理失败: {exc}",
                    code=ErrorCode.IMPORT_FILE_INVALID,
                ) from exc
        finally:
            MetricsService.record_excel_import(
                "total", time.perf_counter() - total_start
            )

    @classmethod
    def _compute_final_status(cls, batch_result: BatchImportResult) -> int:
        """根据批次结果计算最终任务状态"""
        if batch_result.success_count == 0:
            return ImportTaskStatus.FAILED.value
        if batch_result.failed_count > 0:
            return ImportTaskStatus.PARTIAL.value
        return ImportTaskStatus.SUCCESS.value

    @classmethod
    def generate_template(cls, role: int) -> bytes:
        """生成 Excel 导入模板

        Args:
            role: UserRole 枚举值

        Returns:
            Excel 文件二进制内容
        """
        return ExcelTemplateGenerator.generate(role)

    @classmethod
    def generate_error_report(cls, task_id: int) -> bytes:
        """生成错误导出 Excel

        将任务的失败日志导出为 Excel，供管理员修正后重新导入。

        Args:
            task_id: 导入任务 ID

        Returns:
            Excel 文件二进制内容
        """
        task = cls.task_repository.get_by_id(task_id)
        failures = cls.log_repository.list_failures(task_id)

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "失败记录"

        # 表头
        ws.cell(row=1, column=1, value="行号")
        ws.cell(row=1, column=2, value="原始数据")
        ws.cell(row=1, column=3, value="失败原因")

        # 数据行
        for idx, log in enumerate(failures, start=2):
            ws.cell(row=idx, column=1, value=log.row_number)
            ws.cell(row=idx, column=2, value=log.row_data)
            ws.cell(row=idx, column=3, value=log.error_reason or "")

        # 列宽
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 60
        ws.column_dimensions["C"].width = 40

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
