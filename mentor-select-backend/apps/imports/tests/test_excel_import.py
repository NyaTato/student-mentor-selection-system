"""Excel 导入测试

覆盖：
- ExcelTemplateGenerator: 模板生成
- ExcelParser: Excel 解析
- RowValidator: 行校验
- ExcelImporter: 导入器
- ExcelImportService: 完整导入流程
- API 端点: 上传、下载模板、下载错误报告
"""
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from apps.common.enums import (
    ImportRowStatus,
    ImportTaskStatus,
    UserRole,
    UserStatus,
)
from apps.imports.excel.importer import BatchImportResult, ExcelImporter
from apps.imports.excel.parser import ExcelParser
from apps.imports.excel.template import ExcelTemplateGenerator
from apps.imports.excel.validator import RowValidator
from apps.imports.models import ImportLog, ImportTask
from apps.imports.services import ExcelImportService
from apps.students.models import StudentProfile
from apps.teachers.models import TeacherProfile
from apps.users.models import AdminProfile, User


# ===== 辅助函数 =====


def make_student_excel(rows: list[dict]) -> bytes:
    """生成学生导入 Excel

    Args:
        rows: 行数据列表，每行为 {字段名: 值}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "导入模板"

    # 表头（与 constants.STUDENT_COLUMNS 一致）
    headers = [
        "用户名(学号)*",
        "初始密码",
        "姓名*",
        "学号*",
        "专业*",
        "班级*",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)

    # 示例行（第 2 行）
    ws.cell(row=2, column=1, value="20240001")
    ws.cell(row=2, column=3, value="示例")
    ws.cell(row=2, column=4, value="20240001")
    ws.cell(row=2, column=5, value="示例专业")
    ws.cell(row=2, column=6, value="示例班级")

    # 说明行（第 3 行）
    for col in range(1, 7):
        ws.cell(row=3, column=col, value="说明")

    # 数据行（第 4 行开始）
    for row_idx, row_data in enumerate(rows, start=4):
        ws.cell(row=row_idx, column=1, value=row_data.get("username"))
        ws.cell(row=row_idx, column=2, value=row_data.get("password", ""))
        ws.cell(row=row_idx, column=3, value=row_data.get("name"))
        ws.cell(row=row_idx, column=4, value=row_data.get("student_number"))
        ws.cell(row=row_idx, column=5, value=row_data.get("major"))
        ws.cell(row=row_idx, column=6, value=row_data.get("class_name"))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def make_teacher_excel(rows: list[dict]) -> bytes:
    """生成导师导入 Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "导入模板"

    headers = [
        "用户名(工号)*",
        "初始密码",
        "姓名*",
        "工号*",
        "学院*",
        "职称",
        "研究方向",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)

    for col in range(1, 8):
        ws.cell(row=2, column=col, value="示例")
        ws.cell(row=3, column=col, value="说明")

    for row_idx, row_data in enumerate(rows, start=4):
        ws.cell(row=row_idx, column=1, value=row_data.get("username"))
        ws.cell(row=row_idx, column=2, value=row_data.get("password", ""))
        ws.cell(row=row_idx, column=3, value=row_data.get("name"))
        ws.cell(row=row_idx, column=4, value=row_data.get("teacher_number"))
        ws.cell(row=row_idx, column=5, value=row_data.get("college"))
        ws.cell(row=row_idx, column=6, value=row_data.get("title", ""))
        ws.cell(row=row_idx, column=7, value=row_data.get("research_area", ""))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ===== ExcelTemplateGenerator 测试 =====


@pytest.mark.django_db
class TestExcelTemplateGenerator:
    """模板生成器测试"""

    def test_generate_student_template(self):
        """生成学生模板"""
        content = ExcelTemplateGenerator.generate(UserRole.STUDENT.value)
        assert content is not None
        assert len(content) > 0

        # 验证可加载
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        assert ws.title == "导入模板"

        # 验证表头
        headers = [cell.value for cell in ws[1]]
        assert "用户名(学号)*" in headers
        assert "姓名*" in headers
        assert "学号*" in headers

    def test_generate_teacher_template(self):
        """生成导师模板"""
        content = ExcelTemplateGenerator.generate(UserRole.TEACHER.value)
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert "用户名(工号)*" in headers
        assert "工号*" in headers
        assert "学院*" in headers

    def test_generate_admin_template(self):
        """生成管理员模板"""
        content = ExcelTemplateGenerator.generate(UserRole.ADMIN.value)
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert "用户名*" in headers
        assert "姓名*" in headers


# ===== ExcelParser 测试 =====


@pytest.mark.django_db
class TestExcelParser:
    """Excel 解析器测试"""

    def test_parse_student_excel(self):
        """解析学生 Excel"""
        rows_data = [
            {
                "username": "20240001",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机",
                "class_name": "计科2401",
            },
            {
                "username": "20240002",
                "name": "李四",
                "student_number": "20240002",
                "major": "软件",
                "class_name": "软件2401",
            },
        ]
        content = make_student_excel(rows_data)
        rows = ExcelParser.parse(content, UserRole.STUDENT.value)

        assert len(rows) == 2
        assert rows[0]["username"] == "20240001"
        assert rows[0]["name"] == "张三"
        assert rows[0]["student_number"] == "20240001"
        assert rows[0]["major"] == "计算机"
        assert rows[0]["class_name"] == "计科2401"
        assert rows[0]["_row_number"] == 1

    def test_parse_skips_empty_rows(self):
        """跳过空行"""
        rows_data = [
            {"username": "20240001", "name": "张三", "student_number": "20240001",
             "major": "计算机", "class_name": "计科2401"},
        ]
        content = make_student_excel(rows_data)
        # 追加空行
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        ws.cell(row=5, column=1, value=None)
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        rows = ExcelParser.parse(buffer.getvalue(), UserRole.STUDENT.value)
        assert len(rows) == 1

    def test_parse_rejects_oversized_file(self):
        """拒绝超大文件"""
        # 构造超过 5MB 的内容
        content = b"0" * (5 * 1024 * 1024 + 1)
        from apps.common.exceptions import BusinessException

        with pytest.raises(BusinessException) as exc_info:
            ExcelParser.parse(content, UserRole.STUDENT.value)
        assert "超过限制" in str(exc_info.value.detail)


# ===== RowValidator 测试 =====


@pytest.mark.django_db
class TestRowValidator:
    """行校验器测试"""

    def test_validate_valid_student_row(self):
        """校验合法学生行"""
        row_data = {
            "username": "20240001",
            "password": "",
            "name": "张三",
            "student_number": "20240001",
            "major": "计算机",
            "class_name": "计科2401",
        }
        errors = RowValidator.validate(row_data, UserRole.STUDENT.value)
        assert errors == []

    def test_validate_missing_required_field(self):
        """校验必填字段缺失"""
        row_data = {
            "username": "20240001",
            "password": "",
            "name": "",
            "student_number": "20240001",
            "major": "计算机",
            "class_name": "计科2401",
        }
        errors = RowValidator.validate(row_data, UserRole.STUDENT.value)
        assert any("姓名" in e for e in errors)

    def test_validate_invalid_student_number(self):
        """校验学号格式"""
        row_data = {
            "username": "abc123",
            "password": "",
            "name": "张三",
            "student_number": "abc123",
            "major": "计算机",
            "class_name": "计科2401",
        }
        errors = RowValidator.validate(row_data, UserRole.STUDENT.value)
        assert any("数字" in e for e in errors)

    def test_validate_student_number_mismatch(self):
        """学号与用户名不一致"""
        row_data = {
            "username": "20240001",
            "password": "",
            "name": "张三",
            "student_number": "20240002",
            "major": "计算机",
            "class_name": "计科2401",
        }
        errors = RowValidator.validate(row_data, UserRole.STUDENT.value)
        assert any("一致" in e for e in errors)

    def test_validate_valid_teacher_row(self):
        """校验合法导师行"""
        row_data = {
            "username": "T0001",
            "password": "",
            "name": "李教授",
            "teacher_number": "T0001",
            "college": "计算机学院",
            "title": "教授",
            "research_area": "AI",
        }
        errors = RowValidator.validate(row_data, UserRole.TEACHER.value)
        assert errors == []


# ===== ExcelImporter 测试 =====


@pytest.mark.django_db
class TestExcelImporter:
    """导入器测试"""

    def test_import_student_rows(self, academic_year, admin_user):
        """导入学生行"""
        task = ImportTask.objects.create(
            academic_year=academic_year,
            role=UserRole.STUDENT.value,
            status=ImportTaskStatus.PROCESSING.value,
            file_name="test.xlsx",
            file_size=1024,
            started_at=__import__("django").utils.timezone.now(),
            triggered_by=admin_user,
        )

        rows = [
            {
                "username": "20240001",
                "password": "",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机",
                "class_name": "计科2401",
                "_row_number": 1,
            },
            {
                "username": "20240002",
                "password": "",
                "name": "李四",
                "student_number": "20240002",
                "major": "软件",
                "class_name": "软件2401",
                "_row_number": 2,
            },
        ]

        importer = ExcelImporter()
        result = importer.import_rows(task, rows, UserRole.STUDENT.value)

        assert result.success_count == 2
        assert result.failed_count == 0

        # 验证 user 和 profile 创建
        assert User.objects.filter(username="20240001").exists()
        assert StudentProfile.objects.filter(student_number="20240001").exists()

        # 验证 import_log
        assert ImportLog.objects.filter(task=task).count() == 2
        assert ImportLog.objects.filter(
            task=task, status=ImportRowStatus.SUCCESS.value
        ).count() == 2

    def test_import_duplicate_username_skipped(self, academic_year, admin_user):
        """重复 username 跳过"""
        # 预创建一个用户
        User.objects.create_user(
            username="20240001",
            password="TestPass123!",
            role=UserRole.STUDENT.value,
        )

        task = ImportTask.objects.create(
            academic_year=academic_year,
            role=UserRole.STUDENT.value,
            status=ImportTaskStatus.PROCESSING.value,
            file_name="test.xlsx",
            file_size=1024,
            started_at=__import__("django").utils.timezone.now(),
            triggered_by=admin_user,
        )

        rows = [
            {
                "username": "20240001",  # 已存在
                "password": "",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机",
                "class_name": "计科2401",
                "_row_number": 1,
            },
            {
                "username": "20240002",  # 新用户
                "password": "",
                "name": "李四",
                "student_number": "20240002",
                "major": "软件",
                "class_name": "软件2401",
                "_row_number": 2,
            },
        ]

        importer = ExcelImporter()
        result = importer.import_rows(task, rows, UserRole.STUDENT.value)

        # 单行失败不阻塞整批
        assert result.success_count == 1
        assert result.failed_count == 1

        # 验证失败日志记录了错误原因
        failure_log = ImportLog.objects.get(
            task=task, status=ImportRowStatus.FAILED.value
        )
        assert "已存在" in failure_log.error_reason

    def test_import_invalid_row_does_not_block_batch(self, academic_year, admin_user):
        """单行失败不阻塞整批"""
        task = ImportTask.objects.create(
            academic_year=academic_year,
            role=UserRole.STUDENT.value,
            status=ImportTaskStatus.PROCESSING.value,
            file_name="test.xlsx",
            file_size=1024,
            started_at=__import__("django").utils.timezone.now(),
            triggered_by=admin_user,
        )

        rows = [
            {
                "username": "20240001",
                "password": "",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机",
                "class_name": "计科2401",
                "_row_number": 1,
            },
            {
                "username": "abc",  # 学号格式错误
                "password": "",
                "name": "李四",
                "student_number": "abc",
                "major": "软件",
                "class_name": "软件2401",
                "_row_number": 2,
            },
            {
                "username": "20240003",
                "password": "",
                "name": "王五",
                "student_number": "20240003",
                "major": "数学",
                "class_name": "数学2401",
                "_row_number": 3,
            },
        ]

        importer = ExcelImporter()
        result = importer.import_rows(task, rows, UserRole.STUDENT.value)

        assert result.success_count == 2
        assert result.failed_count == 1

    def test_import_teacher_rows(self, academic_year, admin_user):
        """导入导师行"""
        task = ImportTask.objects.create(
            academic_year=academic_year,
            role=UserRole.TEACHER.value,
            status=ImportTaskStatus.PROCESSING.value,
            file_name="test.xlsx",
            file_size=1024,
            started_at=__import__("django").utils.timezone.now(),
            triggered_by=admin_user,
        )

        rows = [
            {
                "username": "T0001",
                "password": "",
                "name": "李教授",
                "teacher_number": "T0001",
                "college": "计算机学院",
                "title": "教授",
                "research_area": "AI",
                "_row_number": 1,
            },
        ]

        importer = ExcelImporter()
        result = importer.import_rows(task, rows, UserRole.TEACHER.value)

        assert result.success_count == 1
        assert TeacherProfile.objects.filter(teacher_number="T0001").exists()


# ===== ExcelImportService 测试 =====


@pytest.mark.django_db
class TestExcelImportService:
    """Excel 导入服务测试"""

    def test_upload_and_import_students(self, academic_year, admin_user):
        """完整学生导入流程"""
        rows_data = [
            {
                "username": "20240001",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机",
                "class_name": "计科2401",
            },
            {
                "username": "20240002",
                "name": "李四",
                "student_number": "20240002",
                "major": "软件",
                "class_name": "软件2401",
            },
        ]
        content = make_student_excel(rows_data)

        task = ExcelImportService.upload_and_import(
            file_content=content,
            file_name="students.xlsx",
            file_size=len(content),
            role=UserRole.STUDENT.value,
            academic_year_id=academic_year.id,
            triggered_by=admin_user,
        )

        assert task.status == ImportTaskStatus.SUCCESS.value
        assert task.total_rows == 2
        assert task.success_count == 2
        assert task.failed_count == 0
        assert task.finished_at is not None

    def test_upload_and_import_with_failures(self, academic_year, admin_user):
        """导入含失败行"""
        rows_data = [
            {
                "username": "20240001",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机",
                "class_name": "计科2401",
            },
            {
                "username": "abc",  # 格式错误
                "name": "李四",
                "student_number": "abc",
                "major": "软件",
                "class_name": "软件2401",
            },
        ]
        content = make_student_excel(rows_data)

        task = ExcelImportService.upload_and_import(
            file_content=content,
            file_name="students.xlsx",
            file_size=len(content),
            role=UserRole.STUDENT.value,
            academic_year_id=academic_year.id,
            triggered_by=admin_user,
        )

        assert task.status == ImportTaskStatus.PARTIAL.value
        assert task.success_count == 1
        assert task.failed_count == 1

    def test_generate_template(self):
        """生成模板"""
        content = ExcelImportService.generate_template(UserRole.STUDENT.value)
        assert len(content) > 0

    def test_generate_error_report(self, academic_year, admin_user):
        """生成错误报告"""
        rows_data = [
            {
                "username": "abc",
                "name": "李四",
                "student_number": "abc",
                "major": "软件",
                "class_name": "软件2401",
            },
        ]
        content = make_student_excel(rows_data)

        task = ExcelImportService.upload_and_import(
            file_content=content,
            file_name="students.xlsx",
            file_size=len(content),
            role=UserRole.STUDENT.value,
            academic_year_id=academic_year.id,
            triggered_by=admin_user,
        )

        report = ExcelImportService.generate_error_report(task.id)
        assert len(report) > 0

        # 验证错误报告内容
        wb = load_workbook(BytesIO(report))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "行号"
        assert ws.cell(row=1, column=3).value == "失败原因"


# ===== API 端点测试 =====


@pytest.mark.django_db
class TestImportAPI:
    """导入 API 端点测试"""

    def test_upload_excel_api(self, authenticated_client, academic_year):
        """上传 Excel API"""
        rows_data = [
            {
                "username": "20240001",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机",
                "class_name": "计科2401",
            },
        ]
        content = make_student_excel(rows_data)

        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile(
            "students.xlsx",
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = authenticated_client.post(
            "/api/v1/imports/upload/",
            {
                "file": upload,
                "role": UserRole.STUDENT.value,
                "academic_year": academic_year.id,
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["code"] == 0
        assert response.data["data"]["success_count"] == 1
        assert response.data["data"]["status"] == ImportTaskStatus.SUCCESS.value

    def test_download_template_api(self, authenticated_client):
        """下载模板 API"""
        response = authenticated_client.get(
            "/api/v1/imports/template/",
            {"role": UserRole.STUDENT.value},
        )

        assert response.status_code == 200
        # RFC 5987 编码的 Content-Disposition
        from urllib.parse import quote
        assert "attachment" in response["Content-Disposition"]
        assert quote("学生导入模板") in response["Content-Disposition"]

    def test_download_error_report_api(self, authenticated_client, academic_year, admin_user):
        """下载错误报告 API"""
        rows_data = [
            {
                "username": "abc",
                "name": "李四",
                "student_number": "abc",
                "major": "软件",
                "class_name": "软件2401",
            },
        ]
        content = make_student_excel(rows_data)

        task = ExcelImportService.upload_and_import(
            file_content=content,
            file_name="students.xlsx",
            file_size=len(content),
            role=UserRole.STUDENT.value,
            academic_year_id=academic_year.id,
            triggered_by=admin_user,
        )

        response = authenticated_client.get(
            f"/api/v1/imports/tasks/{task.id}/error-report/"
        )

        assert response.status_code == 200
        assert "attachment" in response["Content-Disposition"]

    def test_list_tasks_api(self, authenticated_client, academic_year, admin_user):
        """查询任务列表 API"""
        ImportTask.objects.create(
            academic_year=academic_year,
            role=UserRole.STUDENT.value,
            status=ImportTaskStatus.SUCCESS.value,
            file_name="test.xlsx",
            file_size=1024,
            total_rows=1,
            success_count=1,
            triggered_by=admin_user,
        )

        response = authenticated_client.get("/api/v1/imports/tasks/")

        assert response.status_code == 200
        assert response.data["code"] == 0
        assert len(response.data["data"]["list"]) >= 1

    def test_student_cannot_access_imports(self, student_client):
        """学生不能访问导入"""
        response = student_client.get("/api/v1/imports/tasks/")
        assert response.status_code == 403
