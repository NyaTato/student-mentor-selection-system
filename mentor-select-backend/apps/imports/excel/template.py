"""Excel 模板生成器

依据各角色列定义生成带表头和示例数据的 Excel 模板，
供管理员下载后填写并上传导入。

特性：
- 表头加粗、背景色填充
- 必填列表头标注星号
- 第二行为示例数据
- 第三行为字段说明
- 冻结首行
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.imports.excel.constants import get_columns


class ExcelTemplateGenerator:
    """Excel 模板生成器"""

    # 样式常量
    HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    NOTE_FONT = Font(name="微软雅黑", size=10, italic=True, color="808080")
    EXAMPLE_FONT = Font(name="微软雅黑", size=10, color="000000")
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

    @classmethod
    def generate(cls, role: int) -> bytes:
        """生成 Excel 模板文件

        Args:
            role: UserRole 枚举值

        Returns:
            Excel 文件二进制内容
        """
        columns = get_columns(role)
        wb = Workbook()
        ws = wb.active
        ws.title = "导入模板"

        # 第 1 行：表头
        for idx, (field, header, required, max_length, note) in enumerate(
            columns, start=1
        ):
            cell = ws.cell(
                row=1, column=idx, value=f"{header}*" if required else header
            )
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.CENTER

        # 第 2 行：示例数据
        examples = cls._get_examples(role)
        for idx, (field, header, required, max_length, note) in enumerate(
            columns, start=1
        ):
            cell = ws.cell(row=2, column=idx, value=examples.get(field, ""))
            cell.font = cls.EXAMPLE_FONT
            cell.alignment = cls.LEFT

        # 第 3 行：字段说明
        for idx, (field, header, required, max_length, note) in enumerate(
            columns, start=1
        ):
            cell = ws.cell(row=3, column=idx, value=note)
            cell.font = cls.NOTE_FONT
            cell.alignment = cls.LEFT

        # 列宽自适应（基于表头和说明长度）
        for idx, (field, header, required, max_length, note) in enumerate(
            columns, start=1
        ):
            max_len = max(len(header), len(note or ""))
            width = min(max(max_len * 2, 12), 40)
            ws.column_dimensions[get_column_letter(idx)].width = width

        # 冻结首行
        ws.freeze_panes = "A2"

        # 输出为二进制
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def _get_examples(cls, role: int) -> dict:
        """获取各角色的示例数据"""
        from apps.common.enums import UserRole

        if role == UserRole.STUDENT.value:
            return {
                "username": "20240001",
                "password": "",
                "name": "张三",
                "student_number": "20240001",
                "major": "计算机科学与技术",
                "class_name": "计科2401",
            }
        if role == UserRole.TEACHER.value:
            return {
                "username": "T0001",
                "password": "",
                "name": "李教授",
                "teacher_number": "T0001",
                "college": "计算机学院",
                "title": "教授",
                "research_area": "人工智能,机器学习",
            }
        if role == UserRole.ADMIN.value:
            return {
                "username": "admin001",
                "password": "",
                "name": "王管理员",
            }
        return {}
